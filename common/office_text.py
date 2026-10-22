"""Office document text extraction shared by the read and web_fetch tools.

The parsers are optional dependencies, so they are imported lazily (docx) or
supplied by the caller (openpyxl's ``load_workbook``).
"""

from contextlib import contextmanager
from copy import deepcopy


def iter_docx_body_text(document):
    """Yield paragraphs and table rows in body order, including content controls.

    Builds python-docx's block objects directly, which also works on versions
    that predate ``Document.iter_inner_content()``.
    """
    yield from _iter_docx_blocks(document.element.body, document)


def _iter_docx_blocks(container, parent):
    """Traverse paragraphs and tables in either a document body or a cell."""
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    for element in container.iterchildren():
        if element.tag == qn("w:p"):
            fields = list(element.iter(qn("w:fldSimple")))
            if fields:
                # python-docx skips simple-field containers, although their
                # child runs hold the displayed result saved by Word. Unwrap
                # a copy only: read cached text without executing instructions
                # or removing fields from the caller's document.
                element = deepcopy(element)
                for field in list(element.iter(qn("w:fldSimple"))):
                    for child in list(field):
                        field.addprevious(child)
                    field.getparent().remove(field)
            yield Paragraph(element, parent).text
        elif element.tag == qn("w:sdt"):
            content = element.find(qn("w:sdtContent"))
            if content is not None:
                yield from _iter_docx_blocks(content, parent)
        elif element.tag == qn("w:tbl"):
            seen_cells = set()
            for row in Table(element, parent).rows:
                values = []
                for cell in row.cells:
                    # python-docx repeats the same XML cell for every covered
                    # grid slot, including vertical continuations. Keep those
                    # slots empty without dropping equal text in other cells.
                    if cell._tc in seen_cells:
                        values.append("")
                    else:
                        seen_cells.add(cell._tc)
                        values.append("\n".join(_iter_docx_blocks(cell._tc, cell)))
                yield "\t".join(values)


def iter_pptx_shape_text(shapes):
    """Yield slide text, including table cells and nested groups, in shape order."""
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from iter_pptx_shape_text(shape.shapes)
        elif shape.has_table:
            for row in shape.table.rows:
                # A merged cell's origin owns the text; covered cells add none.
                yield "\t".join(cell.text for cell in row.cells if not cell.is_spanned)
        elif shape.has_text_frame:
            for paragraph in shape.text_frame.paragraphs:
                text = paragraph.text.strip()
                if text:
                    yield text

def _sheet_rows(cached_sheet, formula_sheet):
    for cached_row, formula_row in zip(cached_sheet.iter_rows(), formula_sheet.iter_rows()):
        values = []
        for cached, source in zip(cached_row, formula_row):
            value = cached.value
            if value is None and source.data_type == "f":
                # openpyxl's array formulas store their source in .text.
                formula = getattr(source.value, "text", source.value)
                value = f"[Formula: {formula} (not calculated)]"
            values.append(str(value) if value is not None else "")
        yield values


@contextmanager
def spreadsheet_sheets(file_path, load_workbook):
    """Yield (sheet name, rows) pairs and close both read-only views on exit.

    Cached values win, including 0/False. A formula with no cached result is
    shown as its source, marked unevaluated; nothing is recalculated.
    """
    cached = load_workbook(file_path, read_only=True, data_only=True)
    try:
        formulas = load_workbook(file_path, read_only=True, data_only=False)
        try:
            # Some spreadsheet producers leave stale used-range metadata. In
            # read-only mode those bounds silently clip real XML cells, so let
            # both views stream their actual rows/columns before pairing them.
            for sheet in cached.worksheets:
                sheet.reset_dimensions()
                formulas[sheet.title].reset_dimensions()
            yield ((sheet.title, _sheet_rows(sheet, formulas[sheet.title])) for sheet in cached.worksheets)
        finally:
            formulas.close()
    finally:
        cached.close()
