"""Preserve displayed Word field results without evaluating instructions."""
import sys
import tempfile
import unittest
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
pytest.importorskip("docx", reason="Word extraction requires optional python-docx")
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from agent.tools.read.read import Read
from common.office_text import iter_docx_body_text


def add_field(paragraph, value):
    paragraph.add_run("Date: ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), r' DATE \@ "yyyy-MM-dd" ')
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = value
    run.append(text)
    field.append(run)
    paragraph._p.append(field)
    paragraph.add_run(" end")


class WordSimpleFieldsTests(unittest.TestCase):
    def test_public_read_preserves_cached_field_results_in_body_and_cells(self):
        for in_table in (False, True):
            with self.subTest(in_table=in_table), tempfile.TemporaryDirectory() as directory:
                document = Document()
                paragraph = (document.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0]
                             if in_table else document.add_paragraph())
                add_field(paragraph, "2026-10-20")
                path = Path(directory) / "report.docx"
                document.save(path)
                result = Read({"cwd": directory}).execute({"path": str(path)})
                self.assertEqual(result.status, "success", result.result)
                self.assertIn("Date: 2026-10-20 end", result.result["content"])
                self.assertNotIn("yyyy-MM-dd", result.result["content"])

    def test_extraction_is_repeatable_and_leaves_field_xml_unchanged(self):
        document = Document()
        add_field(document.add_paragraph(), "2026-10-20")
        original = document.element.xml
        self.assertEqual(list(iter_docx_body_text(document)), ["Date: 2026-10-20 end"])
        self.assertEqual(list(iter_docx_body_text(document)), ["Date: 2026-10-20 end"])
        self.assertEqual(document.element.xml, original)

    def test_ordinary_paragraph_control(self):
        document = Document()
        document.add_paragraph("ordinary text")
        self.assertEqual(list(iter_docx_body_text(document)), ["ordinary text"])
